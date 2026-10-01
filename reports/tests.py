from decimal import Decimal

from django.test import TestCase
from django.urls import reverse

from core.testing import make_case, make_doctor, make_hospital
from payments.models import Payment


class ReportTests(TestCase):
    def setUp(self):
        self.doctor = make_doctor()
        self.a = make_hospital("Alpha")
        self.b = make_hospital("Beta")
        case = make_case(self.doctor, self.a, fee=1000)
        Payment.objects.create(case=case, amount=Decimal("400"), mode="upi")
        make_case(self.doctor, self.b, fee=2000)
        self.client.force_login(self.doctor.user)

    def test_monthly(self):
        ctx = self.client.get(reverse("reports:monthly")).context
        self.assertEqual(ctx["totals"]["billed"], Decimal("3000"))
        self.assertEqual(ctx["totals"]["received"], Decimal("400"))
        self.assertEqual(ctx["totals"]["outstanding"], Decimal("2600"))
        self.assertEqual(len(ctx["rows"]), 12)

    def test_hospital_wise_sorted_by_outstanding(self):
        rows = self.client.get(reverse("reports:hospitals")).context["rows"]
        self.assertEqual([r["hospital"].name for r in rows], ["Beta", "Alpha"])
        self.assertEqual(rows[1]["outstanding"], Decimal("600"))

    def test_payment_history_and_csv(self):
        ctx = self.client.get(reverse("reports:payments") + "?mode=upi").context
        self.assertEqual(ctx["summary"]["total"], Decimal("400"))
        response = self.client.get(reverse("reports:payments") + "?export=csv")
        self.assertTrue(response["Content-Type"].startswith("text/csv"))
        self.assertIn("Alpha", response.content.decode())


class CasesReportTests(TestCase):
    def setUp(self):
        import datetime

        from django.utils import timezone

        self.today = timezone.localdate()
        self.doctor = make_doctor()
        self.a = make_hospital("Alpha")
        self.b = make_hospital("Beta")
        self.recent = make_case(self.doctor, self.a, days_ago=0, fee=1000, procedure_type="LSCS spinal")
        Payment.objects.create(case=self.recent, amount=Decimal("1000"))
        self.old = make_case(self.doctor, self.b, days_ago=120, fee=2000, patient_reference="IP 55")
        other = make_doctor("dr.x", "Dr. X")
        make_case(other, self.a, fee=99999)  # never visible to self.doctor
        self.client.force_login(self.doctor.user)
        self.url = reverse("reports:cases")
        self.datetime = datetime

    def ids(self, query=""):
        return {c.pk for c in self.client.get(self.url + query).context["page"]}

    def test_default_this_month(self):
        self.assertEqual(self.ids(), {self.recent.pk})

    def test_custom_date_range_hospital_status_and_search(self):
        start = (self.today - self.datetime.timedelta(days=200)).isoformat()
        end = self.today.isoformat()
        self.assertEqual(self.ids(f"?period=custom&start={start}&end={end}"), {self.recent.pk, self.old.pk})
        self.assertEqual(self.ids(f"?period=all&hospital={self.b.pk}"), {self.old.pk})
        self.assertEqual(self.ids("?period=all&status=paid"), {self.recent.pk})
        self.assertEqual(self.ids("?period=all&status=unpaid"), {self.old.pk})
        self.assertEqual(self.ids("?period=all&q=lscs"), {self.recent.pk})
        self.assertEqual(self.ids("?period=all&q=IP 55"), {self.old.pk})

    def test_summary_and_excel_export(self):
        import io

        from openpyxl import load_workbook

        ctx = self.client.get(self.url + "?period=all").context
        self.assertEqual((ctx["summary"]["n"], ctx["summary"]["billed"], ctx["summary"]["outstanding"]),
                         (2, Decimal("3000"), Decimal("2000")))
        response = self.client.get(self.url + "?period=all&export=xlsx")
        self.assertIn("spreadsheetml", response["Content-Type"])
        ws = load_workbook(io.BytesIO(response.content)).active
        values = [[c.value for c in row] for row in ws.iter_rows()]
        self.assertEqual(values[4][:3], ["Case date", "Doctor", "Hospital"])
        self.assertEqual(len(values), 5 + 2 + 1)  # 4 header lines + column row + 2 cases + total
        self.assertEqual(values[-1][0], "TOTAL")
        self.assertNotIn(99999, [v for row in values for v in row])

    def test_csv_export_respects_filters(self):
        response = self.client.get(self.url + f"?period=all&hospital={self.a.pk}&export=csv")
        body = response.content.decode("utf-8-sig")
        self.assertIn("Alpha", body)
        self.assertNotIn("Beta", body)

    def test_last_month_and_fy_presets_parse(self):
        for period in ("last_month", "last_3", "this_fy", "last_fy", "last_12"):
            self.assertEqual(self.client.get(self.url + f"?period={period}").status_code, 200)


class MonthCasesTests(TestCase):
    def setUp(self):
        import datetime

        from django.utils import timezone

        from hospitals.models import Department

        self.doctor = make_doctor()
        self.a, self.b = make_hospital("Alpha"), make_hospital("Beta")
        self.ortho = Department.objects.get(name="Orthopaedics")
        today = timezone.localdate()
        self.month = today.replace(day=1)
        self.paid = make_case(self.doctor, self.a, days_ago=0, fee=1000, department=self.ortho)
        Payment.objects.create(case=self.paid, amount=Decimal("1000"))
        self.unpaid = make_case(self.doctor, self.b, days_ago=0, fee=2000)
        self.last_month = make_case(self.doctor, self.a, days_ago=today.day + 3, fee=500)
        make_case(make_doctor("dr.x", "Dr. X"), self.a, fee=99999)
        self.client.force_login(self.doctor.user)
        self.url = reverse("reports:month_cases", args=[self.month.year, self.month.month])
        self.datetime = datetime

    def ids(self, query=""):
        return {c.pk for c in self.client.get(self.url + query).context["page"]}

    def test_only_that_month_and_own_cases(self):
        response = self.client.get(self.url)
        self.assertContains(response, self.month.strftime("%B %Y"))
        self.assertEqual(self.ids(), {self.paid.pk, self.unpaid.pk})
        self.assertNotContains(response, 'name="period"')
        self.assertNotContains(response, 'name="q"')

    def test_filters(self):
        self.assertEqual(self.ids(f"?hospital={self.b.pk}"), {self.unpaid.pk})
        self.assertEqual(self.ids(f"?department={self.ortho.pk}"), {self.paid.pk})
        self.assertEqual(self.ids("?status=paid"), {self.paid.pk})
        self.assertEqual(self.ids("?status=unpaid"), {self.unpaid.pk})
        # period params from other pages are ignored
        self.assertEqual(self.ids("?period=all&start=2000-01-01"), {self.paid.pk, self.unpaid.pk})

    def test_excel_download(self):
        import io

        from openpyxl import load_workbook

        response = self.client.get(self.url + "?export=xlsx&status=unpaid")
        self.assertIn(f"cases_{self.month:%Y_%m}", response["Content-Disposition"])
        ws = load_workbook(io.BytesIO(response.content)).active
        values = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertIn("Beta", values)
        self.assertNotIn("Alpha", values)

    def test_monthly_summary_links_to_month_page(self):
        response = self.client.get(reverse("reports:monthly"))
        self.assertContains(response, self.url)

    def test_bad_month_404(self):
        self.assertEqual(self.client.get("/reports/monthly/2026/13/").status_code, 404)
