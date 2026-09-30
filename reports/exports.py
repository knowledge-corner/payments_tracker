"""Download report data as Excel (.xlsx) or CSV."""
import csv
import datetime
import io
import re
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _filename(base, ext):
    stamp = timezone.localdate().strftime("%Y-%m-%d")
    return f"{re.sub(r'[^A-Za-z0-9_-]+', '_', base)}_{stamp}.{ext}"


def export_response(fmt, base_name, title, filters_text, columns, rows, totals=None):
    """columns: list of (header, kind) with kind in {"text", "date", "money", "int"}."""
    if fmt == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{_filename(base_name, "csv")}"'
        response.write("﻿")  # BOM so Excel opens ₹ / Indian names correctly
        writer = csv.writer(response)
        writer.writerow([c[0] for c in columns])
        for row in rows:
            writer.writerow([_csv_value(v) for v in row])
        if totals:
            writer.writerow([_csv_value(v) for v in totals])
        return response

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13, color="2F3263")
    ws.append([filters_text])
    ws["A2"].font = Font(italic=True, color="7A7D94")
    ws.append([f"Generated {timezone.localtime():%d %b %Y %H:%M}"])
    ws["A3"].font = Font(color="7A7D94", size=9)
    ws.append([])
    header_row = 5
    ws.append([c[0] for c in columns])
    fill = PatternFill("solid", fgColor="EEF0FA")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="2F3263")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_xlsx_value(v) for v in row])
    if totals:
        ws.append([_xlsx_value(v) for v in totals])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for idx, (header, kind) in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        fmt_code = {"date": "DD/MM/YYYY", "money": "#,##0.00", "int": "0"}.get(kind)
        if fmt_code:
            for r in range(header_row + 1, ws.max_row + 1):
                ws[f"{letter}{r}"].number_format = fmt_code
        width = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows[:500] if idx - 1 < len(r)] + [8])
        ws.column_dimensions[letter].width = min(width + 2, 45)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{max(ws.max_row - (1 if totals else 0), header_row)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{_filename(base_name, "xlsx")}"'
    return response


def _xlsx_value(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def _csv_value(value):
    if isinstance(value, datetime.date):
        return value.strftime("%d/%m/%Y")
    if value is None:
        return ""
    return value
