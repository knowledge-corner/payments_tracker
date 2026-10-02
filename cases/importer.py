"""
Bulk import of completed cases (and optional payments) from Excel.

Flow: build_template() -> doctor fills it -> parse_workbook() validates every
row without touching the database. If any row has a problem nothing is saved
and the rows are shown highlighted; otherwise commit_rows() creates hospitals,
surgeons, cases and payments in one transaction straight away.

Column headers are matched flexibly (e.g. "Date", "Surgery Date" or "Case Date"
all map to the case date) so doctors can also upload their existing sheets.
"""
import datetime
import io
import re
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation

from accounts.models import Doctor
from contacts.models import Surgeon
from hospitals.models import Hospital
from payments.models import Payment

from .models import Case

MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024

# key, template header, required, help text, accepted header aliases
COLUMNS = [
    ("case_date", "Case Date", True, "Date of the case, e.g. 25/09/2026",
     ["date", "case date", "surgery date", "procedure date", "date of surgery", "dos"]),
    ("hospital", "Hospital", True, "Hospital name - pick from the list or type a new one",
     ["hospital", "hospital name", "hospital / clinic", "clinic", "nursing home", "centre", "center"]),
    ("fee", "Fee", True, "Amount billed for the case, e.g. 6000",
     ["fee", "fees", "amount", "charges", "bill amount", "billed", "billed amount", "professional fee"]),
    ("patient_name", "Patient Name", False, "e.g. Sunita Patil",
     ["patient name", "name of patient", "patient"]),
    ("surgeon", "Surgeon", False, "Surgeon's name - a new name is added to your surgeons",
     ["surgeon", "surgeon name", "operating surgeon", "consultant", "called by"]),
    ("procedure_type", "Procedure / Case Type", False, "e.g. LSCS - spinal",
     ["procedure", "procedure type", "case type", "surgery", "operation", "procedure / case type"]),
    ("patient_reference", "Case / Patient Ref", False, "IP no. or bill no.",
     ["patient ref", "case ref", "ip no", "ip number", "uhid", "bill no",
      "reference", "ref", "case / patient ref"]),
    ("notes", "Notes", False, "Anything else", ["notes", "remarks", "comment", "comments"]),
    ("amount_received", "Amount Received", False, "Leave blank if not yet paid",
     ["amount received", "received", "paid", "paid amount", "payment", "payment received", "amount paid"]),
    ("payment_date", "Payment Date", False, "Date the money was received (defaults to case date)",
     ["payment date", "paid on", "received on", "date of payment", "received date"]),
    ("payment_mode", "Payment Mode", False, "UPI, Bank transfer, Cheque, Cash, Card or Other",
     ["payment mode", "mode", "mode of payment", "paid by", "payment method"]),
    ("doctor", "Doctor Username", False, "Admins only: login name of the doctor (e.g. dr.mehta)",
     ["doctor", "doctor username", "anaesthetist", "anesthetist", "username"]),
]
REQUIRED = [c[0] for c in COLUMNS if c[2]]

MODE_ALIASES = {
    "upi": "upi", "gpay": "upi", "google pay": "upi", "phonepe": "upi", "paytm": "upi",
    "bank": "bank", "bank transfer": "bank", "bank transfer neft": "bank", "neft": "bank", "rtgs": "bank",
    "imps": "bank", "online": "bank", "transfer": "bank",
    "cheque": "cheque", "check": "cheque", "chq": "cheque",
    "cash": "cash", "card": "card", "credit card": "card", "debit card": "card", "other": "other",
}
MODE_LABELS = dict(Payment.MODE_CHOICES)


def _key(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


ALIAS_TO_FIELD = {}
for _field, _header, _req, _help, _aliases in COLUMNS:
    for _alias in [_header] + _aliases:
        ALIAS_TO_FIELD.setdefault(_key(_alias), _field)


# --------------------------------------------------------------------------- #
# Template
# --------------------------------------------------------------------------- #

def build_template(include_doctor_column=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    columns = [c for c in COLUMNS if include_doctor_column or c[0] != "doctor"]
    header_fill = PatternFill("solid", fgColor="312E81")
    required_fill = PatternFill("solid", fgColor="B45309")
    for idx, (key, header, required, help_text, _aliases) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=header + (" *" if required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = required_fill if required else header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = 16 if key not in ("hospital", "procedure_type", "notes", "patient_name", "surgeon") else 28
        if key in ("case_date", "payment_date"):
            for r in range(2, 1002):
                ws.cell(row=r, column=idx).number_format = "DD/MM/YYYY"
        if key in ("fee", "amount_received"):
            for r in range(2, 1002):
                ws.cell(row=r, column=idx).number_format = "#,##0"
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # Hospitals list (for the dropdown) - new names are still allowed.
    hs = wb.create_sheet("Hospitals")
    hs.append(["Hospital", "Area / city"])
    for c in hs[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
    hospitals = list(Hospital.objects.active().order_by("name"))
    for h in hospitals:
        hs.append([h.name, h.place])
    hs.column_dimensions["A"].width = 45
    hs.column_dimensions["B"].width = 30

    col_index = {c[0]: i for i, c in enumerate(columns, start=1)}
    if hospitals:
        letter = get_column_letter(col_index["hospital"])
        dv = DataValidation(
            type="list", formula1=f"=Hospitals!$A$2:$A${len(hospitals) + 1}", allow_blank=True,
            showErrorMessage=False,  # allow typing a hospital that is not in the list yet
        )
        dv.add(f"{letter}2:{letter}1001")
        ws.add_data_validation(dv)
    mode_letter = get_column_letter(col_index["payment_mode"])
    mode_dv = DataValidation(type="list", formula1='"UPI,Bank transfer,Cheque,Cash,Card,Other"', allow_blank=True)
    mode_dv.add(f"{mode_letter}2:{mode_letter}1001")
    ws.add_data_validation(mode_dv)

    info = wb.create_sheet("Instructions")
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 70
    info.append(["How to use this template"])
    info["A1"].font = Font(bold=True, size=14, color="312E81")
    info.append([])
    info.append(["1.", "Enter one completed case per row in the 'Cases' sheet (row 2 onwards)."])
    info.append(["2.", "Columns marked * (orange) are required: Case Date, Hospital, Fee."])
    info.append(["3.", "Pick the hospital from the dropdown. A name that is not in the list is added as a new hospital."])
    info.append(["4.", "If the case is already paid (fully or partly), fill Amount Received, Payment Date and Mode."])
    info.append(["5.", "Save the file and upload it (Add Case > Bulk upload). If every row is correct it is saved at once;"])
    info.append(["", "otherwise nothing is saved and the rows with problems are highlighted - fix them and upload again."])
    info.append(["6.", "Uploading the same rows again is safe - exact duplicates are detected and skipped."])
    info.append([])
    info.append(["Column", "What to enter"])
    info["A10"].font = info["B10"].font = Font(bold=True)
    for key, header, required, help_text, _aliases in columns:
        info.append([header + (" *" if required else ""), help_text])
    info.append([])
    info.append(["Example row", "25/09/2026 | Ruby Hall Clinic | 7500 | Sunita Patil | Dr. Kulkarni | LSCS - spinal | IP 45821 | | 7500 | 30/09/2026 | UPI"])
    wb.active = 0

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Parsing & validation (no database writes)
# --------------------------------------------------------------------------- #

@dataclass
class RowResult:
    row_no: int
    values: dict = field(default_factory=dict)  # raw display values
    data: dict = field(default_factory=dict)    # cleaned, JSON-serialisable
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    error_fields: list = field(default_factory=list)  # columns to highlight
    duplicate: bool = False

    def error(self, column, message):
        self.errors.append(message)
        if column not in self.error_fields:
            self.error_fields.append(column)

    @property
    def status(self):
        if self.errors:
            return "error"
        if self.duplicate:
            return "duplicate"
        return "ok"


@dataclass
class ParseResult:
    sheet: str = ""
    header_row: int = 0
    columns: dict = field(default_factory=dict)  # field -> header text found
    rows: list = field(default_factory=list)
    fatal: str = ""

    @property
    def ok_rows(self):
        return [r for r in self.rows if r.status == "ok"]

    def counts(self):
        return {
            "total": len(self.rows),
            "ok": sum(1 for r in self.rows if r.status == "ok"),
            "duplicate": sum(1 for r in self.rows if r.status == "duplicate"),
            "error": sum(1 for r in self.rows if r.status == "error"),
            "new_hospitals": len({r.data.get("hospital_name", "").lower() for r in self.ok_rows if not r.data.get("hospital_id")}),
            "has_errors": any(r.errors for r in self.rows),
            "with_payment": sum(1 for r in self.ok_rows if r.data.get("amount_received")),
        }


DATE_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%y",
                "%d %b %Y", "%d-%b-%Y", "%d-%b-%y", "%d %B %Y", "%b %d %Y", "%d/%b/%Y"]


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial date
        return from_excel(value).date()
    text = str(value).strip().replace(",", " ")
    text = re.sub(r"\s+", " ", text)
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{value}' is not a date (use DD/MM/YYYY)")


def parse_amount(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
    else:
        text = re.sub(r"(?i)rs\.?|inr|₹|,|\s|/-", "", str(value))
        if text in ("", "-"):
            return None
        try:
            amount = Decimal(text)
        except InvalidOperation:
            raise ValueError(f"'{value}' is not a number")
    if amount < 0:
        raise ValueError("Amount cannot be negative")
    return amount.quantize(Decimal("0.01"))


def parse_mode(value):
    key = _key(value)
    if not key:
        return "bank"
    if key in MODE_ALIASES:
        return MODE_ALIASES[key]
    for alias, mode in MODE_ALIASES.items():
        if alias in key:
            return mode
    return "other"


def _display(value):
    if isinstance(value, datetime.datetime):
        value = value.date()
    if isinstance(value, datetime.date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value).strip()


def _find_header(ws):
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        mapping = {}
        for col_idx, cell in enumerate(row):
            name = ALIAS_TO_FIELD.get(_key(str(cell or "").replace("*", "")))
            if name and name not in mapping:
                mapping[name] = (col_idx, str(cell).replace("*", "").strip())
        if {"case_date", "hospital", "fee"} <= set(mapping):
            return row_idx, mapping
    return None, {}


def parse_workbook(uploaded_file, user, default_doctor=None):
    """Read and validate an uploaded .xlsx file. Never writes to the database."""
    result = ParseResult()
    try:
        wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception:
        result.fatal = "Could not read the file. Please upload an Excel .xlsx file (use the template)."
        return result

    ws = wb["Cases"] if "Cases" in wb.sheetnames else wb.worksheets[0]
    result.sheet = ws.title
    header_row, mapping = _find_header(ws)
    if not header_row:
        result.fatal = (
            "Could not find the column headings. The sheet needs at least these columns: "
            "Case Date, Hospital and Fee (download the template to see the layout)."
        )
        return result
    result.header_row = header_row
    result.columns = {k: v[1] for k, v in mapping.items()}

    hospitals = {_key(h.name): h for h in Hospital.objects.active()}
    doctors_by_username = {}
    if user.is_app_admin:
        doctors_by_username = {d.user.username.lower(): d for d in Doctor.objects.select_related("user")}

    today = timezone.localdate()
    seen_keys = {}

    for row_no, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        raw = {name: (row[idx] if idx < len(row) else None) for name, (idx, _h) in mapping.items()}
        if all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in raw.values()):
            continue
        if len(result.rows) >= MAX_ROWS:
            result.fatal = f"The file has more than {MAX_ROWS} rows. Please split it into smaller files."
            break

        r = RowResult(row_no=row_no, values={k: _display(v) for k, v in raw.items()})
        data = r.data

        # Case date
        case_date = None
        try:
            case_date = parse_date(raw.get("case_date"))
            if case_date is None:
                r.error("case_date", "Case date is missing")
            elif case_date > today:
                r.error("case_date", "Case date is in the future")
            else:
                data["case_date"] = case_date.isoformat()
        except ValueError as exc:
            r.error("case_date", f"Case date: {exc}")

        # Hospital
        hospital_name = re.sub(r"\s+", " ", str(raw.get("hospital") or "")).strip()
        if not hospital_name:
            r.error("hospital", "Hospital is missing")
        else:
            match = hospitals.get(_key(hospital_name))
            data["hospital_name"] = match.name if match else hospital_name
            data["hospital_id"] = match.pk if match else None
            if not match:
                r.warnings.append(f"New hospital '{hospital_name}' will be added")

        # Fee
        try:
            fee = parse_amount(raw.get("fee"))
            if fee is None:
                r.error("fee", "Fee is missing")
            else:
                data["fee"] = str(fee)
        except ValueError as exc:
            r.error("fee", f"Fee: {exc}")
            fee = None

        for name, limit in (("procedure_type", 150), ("patient_reference", 100), ("patient_name", 150), ("surgeon", 100)):
            value = _display(raw.get(name))
            if len(value) > limit:
                r.warnings.append(f"{name.replace('_', ' ').capitalize()} shortened to {limit} characters")
                value = value[:limit]
            data[name] = value
        data["notes"] = _display(raw.get("notes"))

        # Optional payment
        try:
            received = parse_amount(raw.get("amount_received"))
        except ValueError as exc:
            r.error("amount_received", f"Amount received: {exc}")
            received = None
        if received:
            if fee is not None and received > fee:
                r.error("amount_received", "Amount received is more than the fee")
            data["amount_received"] = str(received)
            try:
                pay_date = parse_date(raw.get("payment_date")) or (case_date if data.get("case_date") else None)
                if pay_date and pay_date > today:
                    r.error("payment_date", "Payment date is in the future")
                elif pay_date:
                    data["payment_date"] = pay_date.isoformat()
            except ValueError as exc:
                r.error("payment_date", f"Payment date: {exc}")
            data["payment_mode"] = parse_mode(raw.get("payment_mode"))

        # Doctor
        doctor = None
        if user.is_app_admin:
            username = _display(raw.get("doctor")).lower()
            if username:
                doctor = doctors_by_username.get(username)
                if not doctor:
                    r.error("doctor", f"Doctor username '{username}' not found")
            else:
                doctor = default_doctor
                if not doctor:
                    r.error("doctor", "Choose a doctor on the upload form or fill 'Doctor Username'")
        else:
            doctor = user.doctor_profile
        if doctor:
            data["doctor_id"] = doctor.pk
            data["doctor_name"] = str(doctor)

        # Duplicate detection (in the database and within the file)
        if not r.errors:
            dup_key = (data["doctor_id"], _key(data["hospital_name"]), data["case_date"],
                       data["fee"], _key(data["patient_reference"]))
            if dup_key in seen_keys:
                r.warnings.append(f"Same as row {seen_keys[dup_key]} in this file - check it is not entered twice")
            seen_keys.setdefault(dup_key, row_no)
            if data.get("hospital_id"):
                existing = Case.objects.filter(
                    doctor_id=data["doctor_id"], hospital_id=data["hospital_id"],
                    case_date=data["case_date"], fee=Decimal(data["fee"]),
                )
                if data["patient_reference"]:
                    existing = existing.filter(patient_reference__iexact=data["patient_reference"])
                if existing.exists():
                    r.duplicate = True
                    r.warnings.append("Already in the app - skipped")

        result.rows.append(r)

    if not result.rows and not result.fatal:
        result.fatal = "No case rows found under the headings."
    return result


def serialise(result):
    return [asdict(r) for r in result.ok_rows]


COLUMN_HEADERS = [(c[0], c[1]) for c in COLUMNS]


# --------------------------------------------------------------------------- #
# Commit
# --------------------------------------------------------------------------- #

@transaction.atomic
def commit_rows(rows, user, filename):
    """Create hospitals, cases and payments from validated rows. Returns counts."""
    created_cases = created_payments = created_hospitals = created_surgeons = 0
    new_hospitals, surgeons = {}, {}
    label = re.sub(r"[^\w.\- ]+", "", filename or "upload")[:50]
    for row in rows:
        data = row["data"]
        hospital = None
        if data.get("hospital_id"):
            hospital = Hospital.objects.filter(pk=data["hospital_id"]).first()
        if hospital is None:
            key = _key(data["hospital_name"])
            hospital = new_hospitals.get(key) or Hospital.objects.filter(name__iexact=data["hospital_name"]).first()
            if hospital is None:
                hospital = Hospital.objects.create(
                    name=data["hospital_name"], city="", source="doctor", created_by=user,
                    legacy_ref=f"excel:{label}",
                )
                created_hospitals += 1
            new_hospitals[key] = hospital

        surgeon = None
        if data.get("surgeon"):
            key = (data["doctor_id"], _key(data["surgeon"]))
            surgeon = surgeons.get(key) or Surgeon.objects.filter(
                doctor_id=data["doctor_id"], name__iexact=data["surgeon"]).first()
            if surgeon is None:
                surgeon = Surgeon.objects.create(doctor_id=data["doctor_id"], name=data["surgeon"])
                created_surgeons += 1
            surgeons[key] = surgeon
            surgeon.link_hospital(hospital)

        case = Case.objects.create(
            doctor_id=data["doctor_id"], hospital=hospital, surgeon=surgeon,
            patient_name=data.get("patient_name", ""),
            case_date=datetime.date.fromisoformat(data["case_date"]),
            fee=Decimal(data["fee"]),
            procedure_type=data.get("procedure_type", ""),
            patient_reference=data.get("patient_reference", ""),
            notes=data.get("notes", ""),
            source=Case.SOURCE_IMPORT,
            legacy_ref=f"excel:{label}:row{row['row_no']}",
            created_by=user,
        )
        created_cases += 1
        if data.get("amount_received"):
            Payment.objects.create(
                case=case, amount=Decimal(data["amount_received"]),
                payment_date=datetime.date.fromisoformat(data.get("payment_date") or data["case_date"]),
                mode=data.get("payment_mode", "bank"),
                source=Payment.SOURCE_IMPORT,
                legacy_ref=f"excel:{label}:row{row['row_no']}",
                created_by=user,
            )
            created_payments += 1
    return {"cases": created_cases, "payments": created_payments, "hospitals": created_hospitals,
            "surgeons": created_surgeons}
