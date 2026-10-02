import datetime
import io
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.testing import make_admin, make_case, make_doctor, make_hospital
from hospitals.models import Hospital
from payments.models import Payment

from .importer import parse_amount, parse_date, parse_mode
from .models import Case

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_file(rows, headers=None, name="cases.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws.append(headers or ["Case Date *", "Hospital *", "Fee *", "Procedure / Case Type", "Case / Patient Ref",
                          "Notes", "Amount Received", "Payment Date", "Payment Mode"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(), content_type=XLSX)


class ParserUnitTests(TestCase):
    def test_parse_helpers(self):
        self.assertEqual(parse_date("25/09/2026"), datetime.date(2026, 9, 25))
        self.assertEqual(parse_date("2026-09-25"), datetime.date(2026, 9, 25))
        self.assertEqual(parse_date("5-Sep-2026"), datetime.date(2026, 9, 5))
        self.assertEqual(parse_amount("₹ 7,500/-"), Decimal("7500.00"))
        self.assertEqual(parse_amount("Rs. 1,200"), Decimal("1200.00"))
        self.assertEqual(parse_mode("NEFT"), "bank")
        self.assertEqual(parse_mode("Google Pay"), "upi")
        self.assertEqual(parse_mode(""), "bank")
        with self.assertRaises(ValueError):
            parse_date("next monday")


class ImportFlowTests(TestCase):
    def setUp(self):
        self.doctor = make_doctor()
        self.hospital = make_hospital("Ruby Hall Clinic")
        self.client.force_login(self.doctor.user)
        self.day = timezone.localdate() - datetime.timedelta(days=10)

    def upload(self, file, **extra):
        data = {"file": file}
        data.update(extra)
        return self.client.post(reverse("cases:import"), data)

    def test_template_download(self):
        response = self.client.get(reverse("cases:import_template"))
        self.assertEqual(response["Content-Type"], XLSX)
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Cases", "Hospitals", "Instructions"])
        self.assertIn("Ruby Hall Clinic", [c.value for c in wb["Hospitals"]["A"]])
        self.assertNotIn("Doctor Username", [c.value for c in wb["Cases"][1]])

    def test_valid_file_is_saved_straight_away(self):
        headers = ["Case Date *", "Hospital *", "Fee *", "Patient Name", "Surgeon", "Procedure / Case Type",
                   "Case / Patient Ref", "Notes", "Amount Received", "Payment Date", "Payment Mode"]
        f = workbook_file([
            [self.day, "ruby hall clinic", 7500, "Sunita Patil", "Dr. Kulkarni", "LSCS - spinal", "IP 1", "", 7500,
             self.day, "UPI"],
            [self.day.strftime("%d/%m/%Y"), "Brand New Nursing Home", "₹4,000", "", "dr. kulkarni", "GA", "IP 2", "",
             1000, "", "cheque"],
            ["", "", "", "", "", "", "", "", "", "", ""],                       # blank row ignored
        ], headers=headers)
        response = self.upload(f)
        self.assertRedirects(response, reverse("cases:list"))
        self.assertEqual(Case.objects.count(), 2)
        self.assertEqual(Payment.objects.count(), 2)
        self.assertTrue(Hospital.objects.filter(name="Brand New Nursing Home", source="doctor").exists())
        imported = Case.objects.get(patient_reference="IP 1")
        self.assertEqual((imported.hospital, imported.patient_name), (self.hospital, "Sunita Patil"))
        self.assertEqual(imported.source, Case.SOURCE_IMPORT)
        self.assertEqual(imported.payment_status, "paid")
        # one surgeon (name matched case-insensitively), linked to both hospitals
        surgeon = imported.surgeon
        self.assertEqual(Case.objects.get(patient_reference="IP 2").surgeon, surgeon)
        self.assertEqual(surgeon.hospital_links.count(), 2)
        self.assertEqual(Payment.objects.get(case__patient_reference="IP 2").mode, "cheque")

    def test_any_error_saves_nothing_and_highlights_rows(self):
        f = workbook_file([
            [self.day, "Ruby Hall Clinic", 7500, "", "IP 1"],
            [self.day, "Ruby Hall Clinic", "abc", "", ""],                       # bad fee
            [self.day, "Ruby Hall Clinic", 5000, "", "IP 9", "", 6000, "", ""],  # overpaid
        ])
        response = self.upload(f)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Case.objects.count(), 0)
        self.assertEqual([r.row_no for r in response.context["error_rows"]], [3, 4])
        self.assertEqual(response.context["error_rows"][0].error_fields, ["fee"])
        self.assertContains(response, "Nothing was saved")
        self.assertContains(response, "cell-error")

    def test_duplicates_skipped_on_reupload(self):
        make_case(self.doctor, self.hospital, days_ago=10, fee=7500, patient_reference="IP 1")
        self.upload(workbook_file([[self.day, "Ruby Hall Clinic", 7500, "", "IP 1"],
                                   [self.day, "Ruby Hall Clinic", 8000, "", "IP 2"]]))
        self.assertEqual(Case.objects.count(), 2)

    def test_close_hospital_name_becomes_new_hospital(self):
        self.upload(workbook_file([[self.day, "Ruby Hall Clinc", 7500]]))
        self.assertTrue(Hospital.objects.filter(name="Ruby Hall Clinc").exists())

    def test_existing_sheet_with_other_headings(self):
        f = workbook_file(
            [["My practice log"], [], ["Date", "Hospital Name", "Amount", "IP No", "Received", "Mode"],
             [self.day, "Ruby Hall Clinic", 6000, "IP 77", 3000, "NEFT"]],
            headers=["Dr. Test - 2026"],
        )
        self.upload(f)
        case = Case.objects.with_totals().get()
        self.assertEqual((case.patient_reference, case.status), ("IP 77", "partial"))

    def test_rejects_non_excel_file(self):
        response = self.client.post(reverse("cases:import"), {
            "file": SimpleUploadedFile("cases.csv", b"a,b", content_type="text/csv"),
        })
        self.assertFalse(response.context["form"].is_valid())

    def test_admin_import_uses_doctor_column_or_default(self):
        other = make_doctor("dr.rao", "Dr. Rao")
        self.client.force_login(make_admin())
        headers = ["Case Date", "Hospital", "Fee", "Doctor Username"]
        f = workbook_file([[self.day, "Ruby Hall Clinic", 1000, "dr.rao"], [self.day, "Ruby Hall Clinic", 2000, ""]],
                          headers=headers)
        self.upload(f, doctor=self.doctor.pk)
        self.assertEqual(Case.objects.get(fee=1000).doctor, other)
        self.assertEqual(Case.objects.get(fee=2000).doctor, self.doctor)
