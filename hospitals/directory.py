"""
Loading the shared hospital directory.

1. load_starter()      - ~90 well-known Pune / Mumbai hospitals bundled with the app
                         (names and areas only) so search works on day one.
2. import_directory()  - the official "Hospital Directory (National Health Portal)"
                         dataset from data.gov.in (CSV / XLSX), filtered to the
                         districts you choose. Re-running it updates rows instead of
                         creating duplicates.

Source of (2): https://data.gov.in/catalog/hospital-directory-national-health-portal
(Open Government Data Licence - India).
"""
import csv
import io
import json
import re
from pathlib import Path

from django.db import transaction

from .models import Hospital, normalise_name

STARTER_FILE = Path(__file__).parent / "data" / "starter_hospitals.json"


def load_starter(stdout=None):
    created = 0
    for row in json.loads(STARTER_FILE.read_text(encoding="utf-8")):
        _, was_created = Hospital.objects.get_or_create(
            name=row["name"], city=row["city"],
            defaults={
                "area": row.get("area", ""), "ownership": row.get("ownership", ""),
                "aliases": row.get("aliases", ""), "source": "starter", "category": "hospital",
                "district": "Pune" if row["city"] == "Pune" else row["city"],
            },
        )
        created += was_created
    return created


# --------------------------------------------------------------------------- #
# data.gov.in import
# --------------------------------------------------------------------------- #

FIELD_ALIASES = {
    "name": ["hospital_name", "name", "hospital name", "name of hospital", "facility_name", "facility name"],
    "ref": ["sr_no", "srno", "s no", "id", "hospital_id", "hospital_regis_number"],
    "category": ["hospital_category", "category", "facility_type", "type"],
    "care_type": ["hospital_care_type", "care_type"],
    "system": ["discipline_systems_of_medicine", "systems_of_medicine", "system_of_medicine"],
    "address": ["address_original_first_line", "address", "address_first_line", "street"],
    "area": ["location", "locality", "town", "subtown", "subdistrict", "taluka_name", "block_name"],
    "district": ["district", "district_name"],
    "state": ["state", "state_name"],
    "pincode": ["pincode", "pin_code", "pin"],
    "phone": ["telephone", "landline_number", "mobile_number", "emergency_num", "phone", "contact_number"],
    "email": ["hospital_primary_email_id", "email", "email_id"],
    "website": ["website", "web_site"],
}

NON_ALLOPATHIC = ("ayurved", "homoeo", "homeo", "unani", "siddha", "yoga", "naturopathy", "sowa")

CITY_FOR_DISTRICT = {
    "pune": "Pune", "mumbai": "Mumbai", "mumbai suburban": "Mumbai", "mumbai city": "Mumbai",
    "thane": "Thane", "raigad": "Navi Mumbai",
}


def _key(text):
    return re.sub(r"[^a-z0-9]+", "_", str(text or "").strip().lower()).strip("_")


def _rows(uploaded_file):
    name = getattr(uploaded_file, "name", "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(uploaded_file, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        iterator = ws.iter_rows(values_only=True)
        header = [str(c or "") for c in next(iterator)]
        for row in iterator:
            yield dict(zip(header, row))
        return
    raw = uploaded_file.read()
    text = raw.decode("utf-8-sig", errors="replace") if isinstance(raw, bytes) else raw
    yield from csv.DictReader(io.StringIO(text))


def _pick(row, mapping, field):
    for column in mapping.get(field, []):
        value = row.get(column)
        if value not in (None, "") and str(value).strip().upper() not in ("NA", "N/A", "0", "NULL", "-"):
            return str(value).strip()
    return ""


def _category(text):
    t = (text or "").lower()
    if "nursing" in t:
        return "nursing_home"
    if "clinic" in t or "dispensar" in t:
        return "clinic"
    if "day care" in t or "daycare" in t or "surgical" in t:
        return "day_care"
    return "hospital"


def _ownership(text):
    t = (text or "").lower()
    if any(w in t for w in ("govt", "government", "public", "municipal", "corporation")):
        return "government"
    if "trust" in t or "charit" in t:
        return "trust"
    if "private" in t:
        return "private"
    return ""


def _title(text):
    text = re.sub(r"\s+", " ", text or "").strip()
    return text.title() if text.isupper() or text.islower() else text


@transaction.atomic
def import_directory(uploaded_file, districts=("Pune", "Mumbai", "Mumbai Suburban"), allopathic_only=True):
    """Returns a dict of counts. Never touches hospitals added by doctors/admins."""
    wanted = [d.strip().lower() for d in districts if d.strip()]
    counts = {"rows": 0, "matched": 0, "created": 0, "updated": 0, "skipped_system": 0, "skipped_no_name": 0}
    mapping = None
    for row in _rows(uploaded_file):
        counts["rows"] += 1
        if mapping is None:
            keys = {_key(k): k for k in row.keys()}
            mapping = {field: [keys[a] for a in aliases if a in keys] for field, aliases in FIELD_ALIASES.items()}
            if not mapping["name"]:
                raise ValueError("Could not find a hospital name column (e.g. 'Hospital_Name').")
        district = _pick(row, mapping, "district")
        state = _pick(row, mapping, "state")
        area = _pick(row, mapping, "area")
        haystack = f"{district} {area}".lower()
        if wanted and not any(w in haystack for w in wanted):
            continue
        if state and "maharashtra" not in state.lower() and wanted:
            continue
        counts["matched"] += 1
        system = _pick(row, mapping, "system").lower()
        if allopathic_only and system and any(w in system for w in NON_ALLOPATHIC) and "allopath" not in system:
            counts["skipped_system"] += 1
            continue
        name = _title(_pick(row, mapping, "name"))
        if not name:
            counts["skipped_no_name"] += 1
            continue
        city = CITY_FOR_DISTRICT.get(district.lower(), _title(district) or _title(area))
        pincode = re.sub(r"\D", "", _pick(row, mapping, "pincode"))[:6]
        ref = _pick(row, mapping, "ref")
        legacy_ref = f"nhp:{ref}" if ref else f"nhp:{normalise_name(name)}:{pincode}"[:100]
        values = {
            "name": name[:200], "city": city[:80], "district": _title(district)[:80], "area": _title(area)[:120],
            "address": _pick(row, mapping, "address"), "pincode": pincode,
            "phone": _pick(row, mapping, "phone")[:60], "email": _pick(row, mapping, "email")[:254],
            "website": _pick(row, mapping, "website")[:200] if _pick(row, mapping, "website").startswith("http") else "",
            "category": _category(_pick(row, mapping, "category") + " " + _pick(row, mapping, "care_type")),
            "ownership": _ownership(_pick(row, mapping, "category")),
            "source": "directory",
        }
        existing = Hospital.objects.filter(source="directory", legacy_ref=legacy_ref).first()
        if existing is None:
            # Same hospital already present from the starter list? Upgrade it with official details.
            existing = Hospital.objects.filter(source="starter", name__iexact=name, city__iexact=city).first()
        if existing:
            for field, value in values.items():
                if value or field == "source":
                    setattr(existing, field, value)
            existing.legacy_ref = legacy_ref
            existing.save()
            counts["updated"] += 1
        else:
            Hospital.objects.create(legacy_ref=legacy_ref, **values)
            counts["created"] += 1
    return counts


@transaction.atomic
def merge_hospitals(duplicate, keep):
    """Move everything from `duplicate` to `keep`, then hide `duplicate`."""
    from cases.models import Case
    from contacts.models import ContactAffiliation

    moved_cases = Case.objects.filter(hospital=duplicate).update(hospital=keep)
    ContactAffiliation.objects.filter(hospital=duplicate).update(hospital=keep)
    if not keep.phone and duplicate.phone:
        keep.phone = duplicate.phone
    aliases = {a.strip() for a in (keep.aliases + "," + duplicate.name).split(",") if a.strip()}
    keep.aliases = ", ".join(sorted(aliases))[:300]
    keep.save()
    duplicate.merged_into = keep
    duplicate.is_active = False
    duplicate.save(update_fields=["merged_into", "is_active", "updated_at"])
    return moved_cases
