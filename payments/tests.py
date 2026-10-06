import datetime
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from cases.models import Case
from core.testing import make_case, make_doctor, make_hospital

from .models import Payment, PaymentFollowUp
from .services import receivables_for, reminder_state


class PaymentViewTests(TestCase):
    def setUp(self):
        self.doctor = make_doctor()
        self.hospital = make_hospital()
        self.case = make_case(self.doctor, self.hospital, days_ago=3, fee=5000)
        self.client.force_login(self.doctor.user)

    def test_partial_payments(self):
        url = reverse("payments:record_for_case", args=[self.case.pk])
        for amount in ("2000", "1500"):
            self.client.post(url, {"amount": amount, "payment_date": timezone.localdate().isoformat(), "mode": "upi"})
        case = Case.objects.with_totals().get(pk=self.case.pk)
        self.assertEqual(case.total_paid, Decimal("3500"))
        self.assertEqual(case.outstanding, Decimal("1500"))
        self.assertEqual(case.status, "partial")

    def test_overpayment_is_rejected(self):
        response = self.client.post(
            reverse("payments:record_for_case", args=[self.case.pk]),
            {"amount": "6000", "payment_date": timezone.localdate().isoformat(), "mode": "cash"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Payment.objects.exists())

    def test_record_payment_from_unpaid_list(self):
        response = self.client.post(reverse("payments:record"), {
            "case": self.case.pk, "amount": "5000", "payment_date": timezone.localdate().isoformat(), "mode": "bank",
        })
        self.assertRedirects(response, self.case.get_absolute_url())
        self.assertEqual(Case.objects.get(pk=self.case.pk).payment_status, "paid")

    def test_edit_payment_allows_same_amount(self):
        payment = Payment.objects.create(case=self.case, amount=Decimal("5000"))
        response = self.client.post(reverse("payments:edit", args=[payment.pk]), {
            "amount": "5000", "payment_date": timezone.localdate().isoformat(), "mode": "cheque",
        })
        self.assertEqual(response.status_code, 302)
        payment.refresh_from_db()
        self.assertEqual(payment.mode, "cheque")


class ReminderTests(TestCase):
    """Reminders fire only on the expected payment date (default: 30 days) and on dates set in a follow-up."""

    def setUp(self):
        self.doctor = make_doctor()
        self.hospital = make_hospital()
        self.today = timezone.localdate()

    def day(self, n):
        return self.today + datetime.timedelta(days=n)

    def test_no_reminder_before_expected_date(self):
        case = make_case(self.doctor, self.hospital, days_ago=10)  # default expected date = case + 30 days
        state = reminder_state(case)
        self.assertFalse(state.due)
        self.assertEqual(state.next_trigger, case.case_date + datetime.timedelta(days=30))

    def test_old_7_14_21_day_reminders_are_gone(self):
        for age in (7, 8, 14, 21, 28, 29):
            case = make_case(self.doctor, self.hospital, days_ago=age)
            self.assertFalse(reminder_state(case).due, age)

    def test_due_on_default_30_days(self):
        self.assertTrue(reminder_state(make_case(self.doctor, self.hospital, days_ago=30)).due)

    def test_due_on_expected_date_that_was_entered(self):
        case = make_case(self.doctor, self.hospital, days_ago=1, due_date=self.day(1))
        self.assertFalse(reminder_state(case).due)
        self.assertTrue(reminder_state(case, today=self.day(1)).due)  # on the expected date itself
        self.assertTrue(reminder_state(case, today=self.day(5)).due)  # still due until followed up

    def test_followup_clears_and_remind_again_date_raises_it(self):
        case = make_case(self.doctor, self.hospital, days_ago=31)
        self.assertTrue(reminder_state(case).due)
        self.client.force_login(self.doctor.user)
        self.client.post(reverse("payments:followup_add", args=[case.pk]), {
            "followup_date": self.today.isoformat(), "method": "call", "next_reminder": self.day(4).isoformat(),
        })
        case = Case.objects.get(pk=case.pk)
        self.assertEqual(case.followup_snoozed_until, self.day(4))
        self.assertFalse(reminder_state(case).due)
        self.assertEqual(reminder_state(case).next_trigger, self.day(4))
        self.assertTrue(reminder_state(case, today=self.day(4)).due)

    def test_followup_without_date_means_no_more_reminders(self):
        case = make_case(self.doctor, self.hospital, days_ago=31)
        PaymentFollowUp.objects.create(case=case, followup_date=self.today)
        state = reminder_state(case, today=self.day(60))
        self.assertFalse(state.due)
        self.assertIsNone(state.next_trigger)

    def test_followup_before_expected_date_keeps_that_reminder(self):
        case = make_case(self.doctor, self.hospital, days_ago=2, due_date=self.day(3))
        PaymentFollowUp.objects.create(case=case, followup_date=self.today)
        self.assertTrue(reminder_state(case, today=self.day(3)).due)

    def test_remind_again_date_must_be_in_future(self):
        case = make_case(self.doctor, self.hospital, days_ago=31)
        self.client.force_login(self.doctor.user)
        response = self.client.post(reverse("payments:followup_add", args=[case.pk]), {
            "followup_date": self.today.isoformat(), "method": "call", "next_reminder": self.today.isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("next_reminder", response.context["form"].errors)

    def test_remind_me_in_days(self):
        case = make_case(self.doctor, self.hospital, days_ago=31)
        self.client.force_login(self.doctor.user)
        self.client.post(reverse("payments:followup_snooze", args=[case.pk]), {"days": "3"})
        case.refresh_from_db()
        self.assertEqual(case.followup_snoozed_until, self.day(3))
        self.assertFalse(reminder_state(case).due)
        self.assertTrue(reminder_state(case, today=self.day(3)).due)

    def test_promised_date_becomes_the_reminder(self):
        case = make_case(self.doctor, self.hospital, days_ago=31)
        self.client.force_login(self.doctor.user)
        self.client.post(reverse("payments:followup_add", args=[case.pk]), {
            "followup_date": self.today.isoformat(), "method": "call",
            "promised_payment_date": self.day(5).isoformat(), "notes": "Will pay Friday",
        })
        case.refresh_from_db()
        self.assertEqual(case.followup_snoozed_until, self.day(5))
        self.assertTrue(reminder_state(case, today=self.day(5)).due)

    def test_quick_mark_followed_up(self):
        case = make_case(self.doctor, self.hospital, days_ago=31)
        self.assertTrue(reminder_state(case).due)
        self.client.force_login(self.doctor.user)
        self.client.post(reverse("payments:followup_quick", args=[case.pk]))
        self.assertEqual(case.followups.count(), 1)
        self.assertFalse(reminder_state(Case.objects.get(pk=case.pk)).due)

    def test_paid_cases_never_need_followup(self):
        case = make_case(self.doctor, self.hospital, days_ago=50, fee=1000)
        Payment.objects.create(case=case, amount=Decimal("1000"))
        self.assertEqual(receivables_for(self.doctor.user), [])

    def test_receivables_sorted_oldest_first(self):
        newer = make_case(self.doctor, self.hospital, days_ago=5)
        older = make_case(self.doctor, self.hospital, days_ago=50)
        self.client.force_login(self.doctor.user)
        response = self.client.get(reverse("receivables:list"))
        self.assertEqual([c.pk for c in response.context["cases"]], [older.pk, newer.pk])
        response = self.client.get(reverse("receivables:list") + "?view=overdue")
        self.assertEqual([c.pk for c in response.context["cases"]], [older.pk])


class ReceivedByTests(TestCase):
    def setUp(self):
        self.doctor = make_doctor()
        self.case = make_case(self.doctor, make_hospital(), fee=5000)
        self.client.force_login(self.doctor.user)

    def test_record_payment_with_received_by(self):
        url = reverse("payments:record_for_case", args=[self.case.pk])
        self.assertContains(self.client.get(url), 'name="received_by"')
        self.client.post(url, {"amount": "2000", "payment_date": timezone.localdate().isoformat(), "mode": "upi",
                               "received_by": "Clinic reception"})
        payment = Payment.objects.get()
        self.assertEqual(payment.received_by, "Clinic reception")
        self.assertContains(self.client.get(self.case.get_absolute_url()), "received by Clinic reception")
        # offered as a suggestion next time
        self.assertContains(self.client.get(url), '<option value="Clinic reception">')

    def test_received_by_is_optional_and_in_excel(self):
        import io

        from openpyxl import load_workbook

        url = reverse("payments:record_for_case", args=[self.case.pk])
        self.client.post(url, {"amount": "1000", "payment_date": timezone.localdate().isoformat(), "mode": "cash"})
        Payment.objects.create(case=self.case, amount=Decimal("500"), received_by="Dr. Mehta")
        self.assertEqual(Payment.objects.count(), 2)
        response = self.client.get(reverse("reports:payments") + "?period=all&export=xlsx")
        values = [c.value for row in load_workbook(io.BytesIO(response.content)).active.iter_rows() for c in row]
        self.assertIn("Received by", values)
        self.assertIn("Dr. Mehta", values)
